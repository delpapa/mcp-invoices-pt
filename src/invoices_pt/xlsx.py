"""Local workbook backend (.xlsx, opens in Excel, LibreOffice or Numbers). The default.

One workbook per tax year with an Expenses and an Income sheet in the same layout as the Google Sheet. The cells
keep the formulas so the workbook stays readable and editable; the tools compute the numbers with `bookmath`, since a
file written by Python carries no calculated values (Excel and LibreOffice calculate them when you open it).
Filed quarters are recorded in a hidden `_meta` sheet; their declaration values are stored as fixed numbers.
"""

from __future__ import annotations

import datetime as dt
import os
import re
import tempfile

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from . import bookmath, columns, layout
from .config import Config

END_ROW = 5000          # formulas and validation cover the table down to this row
META = "_meta"
NUMBER_FORMATS = {"date": "dd.mm.yyyy", "num": "#,##0.00", "fx": "0.0000", "pct": "0%"}
FILL = {"box": PatternFill("solid", fgColor="FFFBEB"), "irs": PatternFill("solid", fgColor="EDF7ED"),
        "head": PatternFill("solid", fgColor="EDEDED"), "block": PatternFill("solid", fgColor="FFF2BF"),
        "filed": PatternFill("solid", fgColor="EDEDED")}


def _spec(kind: str) -> list[tuple]:
    return layout.EXPENSE_COLUMNS if kind == "expenses" else layout.INCOME_COLUMNS


class WorkbookOpen(RuntimeError):
    pass


class XlsxBook:
    def __init__(self, cfg: Config, year: int):
        self.cfg, self.year = cfg, year
        self.path = cfg.workbook(year)
        self.names = {k: cfg.sheet_name(k, year) for k in ("expenses", "income")}
        self.expenses = XlsxExpenses(self)
        self.income = XlsxIncome(self)

    def load(self, data_only: bool = False) -> openpyxl.Workbook:
        if not os.path.exists(self.path):
            raise FileNotFoundError(f"{self.path} does not exist; run create_year_sheet({self.year}) first")
        return openpyxl.load_workbook(self.path, data_only=data_only)

    def check_closed(self) -> None:
        folder, name = os.path.split(self.path)
        for lock in (f".~lock.{name}#", f"~${name}"):    # LibreOffice, Excel
            if os.path.exists(os.path.join(folder, lock)):
                raise WorkbookOpen(f"{name} is open in a spreadsheet program; close it and try again")

    def save(self, wb: openpyxl.Workbook) -> None:
        self.check_closed()
        wb.calculation.fullCalcOnLoad = True           # Excel / LibreOffice recalculate on open
        fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=os.path.dirname(self.path) or ".")
        os.close(fd)
        wb.save(tmp)
        os.replace(tmp, self.path)                     # atomic: a crash never leaves half a workbook

    @staticmethod
    def filed(wb: openpyxl.Workbook) -> dict[str, str]:
        if META not in wb.sheetnames:
            return {}
        return {r[1]: r[2] for r in wb[META].iter_rows(values_only=True) if r and r[0] == "filed" and r[1]}


class XlsxTable:
    KIND = ""

    def __init__(self, book: XlsxBook):
        self.book = book
        self.columns = {**columns.fields(self.KIND), **book.cfg.column_overrides(self.KIND)}
        self.date_field = next(iter(columns.fields(self.KIND)))

    # ---- reading
    def _sheet(self, wb):
        return wb[self.book.names[self.KIND]]

    def _header(self, ws) -> tuple[int, dict[str, int]]:
        first = self.columns[self.date_field]
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200)):
            if row[0].value is not None and str(row[0].value).strip() == first:
                index = {str(c.value).strip(): c.column for c in row if c.value is not None}
                missing = [h for h in self.columns.values() if h.strip() not in index]
                if missing:
                    raise KeyError(f"{ws.title}: headers not found: {missing}; map them in the config")
                return row[0].row, index
        raise KeyError(f"{ws.title}: no header row starting with {first!r}")

    def _records(self, wb) -> list[dict]:
        ws = self._sheet(wb)
        head, idx = self._header(ws)
        out = []
        for r in range(head + 1, ws.max_row + 1):
            raw = {f: ws.cell(r, idx[h.strip()]).value for f, h in self.columns.items()}
            if all(v in (None, "") for v in raw.values()):
                continue
            rec = {}
            for f, v in raw.items():
                if f in columns.DATE_FIELDS:
                    d = bookmath.day(v)
                    rec[f] = d.strftime("%d.%m.%Y") if d else ("" if v is None else str(v))
                elif f in columns.NUMBER_FIELDS:
                    rec[f] = "" if v is None else v
                else:
                    rec[f] = bookmath.text_id(v)
            rec["row"] = r
            rec["quarter"] = bookmath.quarter(rec[self.date_field])
            out.append(rec)
        return out

    def records(self) -> list[dict]:
        return self._records(self.book.load())

    def filed_quarters(self) -> set[str]:
        return set(XlsxBook.filed(self.book.load()))

    def find(self, value: str, field: str = "invoice_no", supplier_vat: str = "") -> int | None:
        norm = lambda s: re.sub(r"\s+", "", str(s or "")).upper()
        for rec in self.records():
            if norm(value) and norm(rec.get(field)) == norm(value) and (
                    not supplier_vat or norm(rec.get("supplier_vat")) in ("", norm(supplier_vat))):
                return rec["row"]
        return None

    # ---- writing
    def _cell_value(self, field: str, v):
        if v in (None, ""):
            return None
        if field in columns.DATE_FIELDS:
            return bookmath.day(v) or v
        if field in columns.NUMBER_FIELDS:
            x = bookmath.num(v)
            return v if x is None else x
        return str(v)

    def add(self, row: dict, dry_run: bool = True, allow_filed: bool = False) -> dict:
        day = bookmath.day(row[self.date_field])
        if not day:
            return {"error": f"{self.date_field} must be a date (YYYY-MM-DD)"}
        q = bookmath.quarter(day)
        wb = self.book.load()
        if q in XlsxBook.filed(wb) and not allow_filed:
            return {"error": f"{q} is filed and locked; pass allow_filed only for a deliberate correction"}
        recs = self._records(wb)
        for field in columns.UNIQUE[self.KIND]:
            norm = lambda s: re.sub(r"\s+", "", str(s or "")).upper()
            if row.get(field) and any(norm(r.get(field)) == norm(row[field]) for r in recs):
                dup = next(r["row"] for r in recs if norm(r.get(field)) == norm(row[field]))
                return {"error": f"{field} {row[field]} is already in row {dup}"}
        cells = {self.columns[k]: v for k, v in row.items() if k in self.columns and v not in (None, "")}
        if dry_run:
            return {"dry_run": True, "workbook": self.book.path, "sheet": self.book.names[self.KIND], "quarter": q,
                    "cells": cells}
        ws = self._sheet(wb)
        head, idx = self._header(ws)
        formula_headers = set(layout.row_formulas(self.KIND, 0))
        width = max(idx.values())
        # every existing row as {column: value}, keeping columns the user added by hand
        rows = []
        for r in range(head + 1, ws.max_row + 1):
            vals = {c: ws.cell(r, c).value for c in range(1, width + 1)}
            inputs = [vals[idx[h.strip()]] for h in self.columns.values()]
            if any(v not in (None, "") for v in inputs):
                rows.append(vals)
        new = {c: None for c in range(1, width + 1)}
        for field, header in self.columns.items():
            new[idx[header.strip()]] = self._cell_value(field, row.get(field))
        new["_new"] = True
        rows.append(new)
        date_col = idx[self.columns[self.date_field].strip()]
        rows.sort(key=lambda v: (bookmath.day(v[date_col]) is None, bookmath.day(v[date_col]) or dt.date.max))
        for r in range(head + 1, max(ws.max_row, head + len(rows)) + 1):     # clear, then rewrite sorted
            for c in range(1, width + 1):
                ws.cell(r, c).value = None
        header_by_col = {c: h for h, c in idx.items()}
        spec = {h: fmt for h, _, fmt in _spec(self.KIND)}
        new_row = None
        for k, vals in enumerate(rows):
            r = head + 1 + k
            formulas = layout.row_formulas(self.KIND, r)
            for c in range(1, width + 1):
                h = header_by_col.get(c, "")
                cell = ws.cell(r, c)
                cell.value = formulas[h] if h in formula_headers else vals.get(c)
                if spec.get(h):
                    cell.number_format = NUMBER_FORMATS[spec[h]]
                if h.startswith("Box ") or h == "Net VAT benefit €":
                    cell.fill = FILL["box"]
                elif h.startswith("IRS deductible"):
                    cell.fill = FILL["irs"]
            if vals.get("_new"):
                new_row = r
        self.book.save(wb)
        return {"written": True, "workbook": self.book.path, "sheet": ws.title, "quarter": q, "row": new_row}


class XlsxIncome(XlsxTable):
    KIND = "income"

    def summary(self) -> list[dict]:
        wb = self.book.load()
        return bookmath.income_totals(self.book.expenses._records(wb), self._records(wb), self.book.year)


class XlsxExpenses(XlsxTable):
    KIND = "expenses"

    def _block(self, ws) -> tuple[dict[int, int], dict[str, int]]:
        """Row of each 'Box n' line above the table, and the column of each quarter."""
        head, _ = self._header(ws)
        boxes, quarters = {}, {}
        for r in range(1, head):
            v = ws.cell(r, 1).value
            if v and (m := re.match(r"^Box (\d+)$", str(v).strip())):
                boxes[int(m.group(1))] = r
            if not quarters:
                found = {str(ws.cell(r, c).value).strip(): c for c in range(1, 12) if ws.cell(r, c).value in bookmath.QUARTERS}
                if len(found) == 4:
                    quarters = found
        if not boxes or not quarters:
            raise KeyError("no Declaração Periódica block above the table")
        return boxes, quarters

    def _declaration(self, wb) -> tuple[dict[str, dict[int, float]], dict[str, str]]:
        ws = self._sheet(wb)
        boxes, quarters = self._block(ws)
        filed = XlsxBook.filed(wb)
        frozen = {q: {b: bookmath.num(ws.cell(r, quarters[q]).value) or 0.0 for b, r in boxes.items()} for q in filed}
        opening = bookmath.num(ws.cell(boxes[61], quarters["Q1"]).value) or 0.0
        all_q = bookmath.declaration(self._records(wb), self.book.income._records(wb), opening, frozen)
        return all_q, filed

    def declaration(self, q: str) -> dict:
        wb = self.book.load()
        all_q, filed = self._declaration(wb)
        ws = self._sheet(wb)
        boxes, _ = self._block(ws)
        rows = [{"box": b, "label": ws.cell(r, 3).value or "", "value": all_q[q].get(b)} for b, r in
                sorted(boxes.items(), key=lambda kv: kv[1])]
        return {"quarter": q, "boxes": rows, "filed": q in filed, "live_formulas": q not in filed}

    def set_declaration(self, q: str, box_values: dict[int, float]) -> None:
        wb = self.book.load()
        ws = self._sheet(wb)
        boxes, quarters = self._block(ws)
        for b, v in box_values.items():
            if int(b) in boxes:
                ws.cell(boxes[int(b)], quarters[q]).value = v
        self.book.save(wb)

    def freeze(self, q: str, dry_run: bool = True) -> dict:
        wb = self.book.load()
        all_q, filed = self._declaration(wb)
        exp_rows = [r["row"] for r in self._records(wb) if r["quarter"] == q]
        inc_rows = [r["row"] for r in self.book.income._records(wb) if r["quarter"] == q]
        span = lambda rows: f"{min(rows)}-{max(rows)}" if rows else "none"
        plan = {"quarter": q, "expense_rows": span(exp_rows), "income_rows": span(inc_rows),
                "already_filed": q in filed}
        if dry_run:
            return {"dry_run": True, **plan}
        ws = self._sheet(wb)
        boxes, quarters = self._block(ws)
        for b, r in boxes.items():
            cell = ws.cell(r, quarters[q])
            cell.value = all_q[q].get(b, 0.0)
            cell.fill = FILL["filed"]
        if META not in wb.sheetnames:
            wb.create_sheet(META).sheet_state = "hidden"
        if q not in filed:
            wb[META].append(["filed", q, dt.datetime.now().isoformat(timespec="seconds")])
        self.book.save(wb)
        return {"frozen": True, **plan}


def create_workbook(path: str, year: int, expenses_name: str, income_name: str, opening_credit: float = 0.0) -> dict:
    """A new workbook for a tax year, in the same layout as the Google Sheet. Refuses to overwrite."""
    if os.path.exists(path):
        raise FileExistsError(f"{path} already exists")
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    wb = openpyxl.Workbook()
    exp = wb.active
    exp.title = expenses_name
    inc = wb.create_sheet(income_name)
    wb.create_sheet(META).sheet_state = "hidden"
    for ws, rows in ((exp, layout.expense_values(year, income_name, opening_credit, end=END_ROW)),
                     (inc, layout.income_values(year, expenses_name, end=END_ROW))):
        for r, line in enumerate(rows, 1):
            for c, v in enumerate(line, 1):
                if v not in (None, ""):
                    ws.cell(r, c).value = v
    bold = Font(bold=True)
    for ws, kind, header_row in ((exp, "expenses", layout.EXP_HEADER_ROW), (inc, "income", layout.INC_HEADER_ROW)):
        ws["A1"].font = Font(bold=True, size=12)
        for c, (h, _, _) in enumerate(_spec(kind), 1):
            cell = ws.cell(header_row, c)
            cell.font, cell.fill = bold, FILL["head"]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            ws.column_dimensions[layout.col_letter(c - 1)].width = 14 if h.startswith("Box") else 18
        ws.row_dimensions[header_row].height = 32
    for c in range(1, 9):
        exp.cell(layout.BLOCK_TITLE_ROW, c).font = bold
        exp.cell(layout.BLOCK_TITLE_ROW, c).fill = FILL["block"]
    block_end = layout.BLOCK_TITLE_ROW + len(layout.BLOCK) + len(layout.RESULTS)
    for r in range(layout.BLOCK_TITLE_ROW + 1, block_end + 1):
        exp.cell(r, 1).font = bold
        for c in range(5, 9):
            exp.cell(r, c).number_format = NUMBER_FORMATS["num"]
    for r in range(layout.SUMS_ROW, layout.SUMS_ROW + 6):
        exp.cell(r, 1).font = bold
        for c in range(13, len(layout.EXPENSE_COLUMNS) + 1):
            exp.cell(r, c).number_format = NUMBER_FORMATS["num"]
            if r == layout.SUMS_ROW:
                exp.cell(r, c).font = bold
                exp.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    exp.row_dimensions[layout.SUMS_ROW].height = 32
    inc.cell(2, 1).font = bold
    for r in range(3, 6):
        for c in range(3, 8):
            inc.cell(r, c).number_format = NUMBER_FORMATS["num"]
    exp.column_dimensions["C"].width = 44
    inc.column_dimensions["A"].width = 30
    exp.freeze_panes = "D1"
    names = [h for h, _, _ in layout.EXPENSE_COLUMNS]
    inc_names = [h for h, _, _ in layout.INCOME_COLUMNS]
    first, end = layout.EXP_FIRST_ROW, END_ROW
    for ws, col, values in ((exp, names.index("VAT code"), list(layout.CODES)),
                            (exp, names.index("Document"), layout.DOCUMENT_LABELS),
                            (inc, inc_names.index("VAT case"), list(layout.SALES_CODES))):
        letter = layout.col_letter(col)
        start = first if ws is exp else layout.INC_FIRST_ROW
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(f"{letter}{start}:{letter}{end}")
    doc = layout.col_letter(names.index("Document"))
    exp.conditional_formatting.add(f"{doc}{first}:{doc}{end}", FormulaRule(
        formula=[f'ISNUMBER(SEARCH("MISSING",{doc}{first}))'], font=Font(bold=True, color="CC1A1A")))
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    return {"workbook": path, "expenses_sheet": expenses_name, "income_sheet": income_name,
            "opening_credit": opening_credit}
