"""The local workbook backend, end to end and offline. If LibreOffice is installed, the workbook is also
recalculated by it, to check that the sheet formulas give the same numbers the tools compute in Python."""

import os
import shutil
import subprocess

import openpyxl
import pytest

from invoices_pt import bookmath, layout
from invoices_pt.books import create_year, open_books, opening_credit
from invoices_pt.config import Config
from invoices_pt.xlsx import WorkbookOpen

from test_live_layout import EXPECTED, PURCHASES, SALES


@pytest.fixture
def books(tmp_path):
    cfg = Config({"taxpayer": {"nif": "999999990"}, "paths": {"year_root": str(tmp_path / "{year}")}})
    create_year(cfg, 2099, 100.0)
    exp, inc = open_books(cfg, 2099)
    for row in PURCHASES:
        assert exp.add(row, dry_run=False)["written"], row
    for row in SALES:
        assert inc.add(row, dry_run=False)["written"], row
    return cfg, exp, inc


def boxes(exp, q):
    return {b["box"]: b["value"] for b in exp.declaration(q)["boxes"]}


def test_workbook_is_created_where_configured(tmp_path):
    cfg = Config({"paths": {"year_root": str(tmp_path / "{year}")}})
    out = create_year(cfg, 2099, 12.5)
    assert out["workbook"] == str(tmp_path / "2099" / "Bookkeeping_2099.xlsx") and os.path.exists(out["workbook"])
    with pytest.raises(FileExistsError):
        create_year(cfg, 2099, 0)
    wb = openpyxl.load_workbook(out["workbook"])
    assert wb.sheetnames[:2] == ["Expenses", "Income"]
    assert wb["Expenses"].cell(layout.EXP_HEADER_ROW, 1).value == "Date"


def test_boxes_and_totals(books):
    _, exp, inc = books
    for q, expected in EXPECTED.items():
        got = boxes(exp, q)
        assert all(abs(got[b] - v) < 0.005 for b, v in expected.items()), (q, got)
    totals = {t["item"][:12]: t for t in inc.summary()}
    assert totals["Gross income"]["Q1"] == 1500.0
    assert (totals["Income € by "]["Q1"], totals["Income € by "]["Q2"]) == (500.0, 1000.0)
    assert [r["quarter"] for r in exp.records()] == ["Q1", "Q1", "Q1", "Q2"]


def test_rows_sorted_duplicates_refused(books):
    _, exp, _ = books
    late = dict(PURCHASES[0], date="2099-01-02", invoice_no="A0")
    assert exp.add(late, dry_run=False)["row"] == layout.EXP_FIRST_ROW        # sorted in before the others
    assert "already in row" in exp.add(dict(PURCHASES[0]), dry_run=True)["error"]
    rec = exp.records()[0]
    assert (rec["invoice_no"], rec["amount"], rec["code"]) == ("A0", 20, "NON-EU-RC")


def test_freeze_locks_quarter_and_keeps_chain(books):
    cfg, exp, inc = books
    assert exp.freeze("Q1", dry_run=False)["frozen"]
    assert exp.filed_quarters() == {"Q1"} and inc.filed_quarters() == {"Q1"}
    assert "locked" in exp.add(dict(PURCHASES[0], invoice_no="late"), dry_run=True)["error"]
    assert exp.declaration("Q1")["filed"] and boxes(exp, "Q1")[96] == 0.0
    # a later correction to a Q1 row (e.g. by hand) no longer changes the filed values
    wb = openpyxl.load_workbook(exp.book.path)
    wb["Expenses"].cell(layout.EXP_FIRST_ROW, 10).value = 999
    wb.save(exp.book.path)
    assert boxes(exp, "Q1")[3] == EXPECTED["Q1"][3]
    assert opening_credit(cfg, 2100) == boxes(exp, "Q4")[96]


def test_refuses_to_write_an_open_workbook(books):
    _, exp, _ = books
    folder, name = os.path.split(exp.book.path)
    open(os.path.join(folder, f".~lock.{name}#"), "w").close()
    with pytest.raises(WorkbookOpen):
        exp.add(dict(PURCHASES[0], invoice_no="B1"), dry_run=False)


@pytest.mark.skipif(not shutil.which("soffice"), reason="LibreOffice not installed")
def test_formulas_agree_with_python(books, tmp_path):
    """Let LibreOffice calculate the workbook and compare its numbers with the tools' (the Python mirror)."""
    _, exp, inc = books
    out = tmp_path / "recalculated"
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), exp.book.path],
                   check=True, capture_output=True, timeout=180)
    wb = openpyxl.load_workbook(out / os.path.basename(exp.book.path), data_only=True)
    ws = wb["Expenses"]
    rows = {ws.cell(r, 1).value: r for r in range(1, layout.EXP_HEADER_ROW) if ws.cell(r, 1).value}
    for qi, q in enumerate(bookmath.QUARTERS):
        mine = boxes(exp, q)
        for box, value in mine.items():
            sheet = ws.cell(rows[f"Box {box}"], 5 + qi).value
            assert abs((sheet or 0) - value) < 0.005, (q, box, sheet, value)
    irs_col = [h for h, _, _ in layout.EXPENSE_COLUMNS].index("IRS deductible € (Anexo B)") + 1
    mine_irs = {t["item"][:8]: t for t in inc.summary()}["Business"]
    assert abs(ws.cell(layout.SUMS_ROW + 5, irs_col).value - mine_irs["2099"]) < 0.005
    totals = wb["Income"]
    assert abs(totals.cell(3, 3).value - 1500.0) < 0.005 and abs(totals.cell(4, 4).value - 1000.0) < 0.005
